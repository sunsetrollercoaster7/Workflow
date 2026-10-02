#!/usr/bin/env python3
"""Convert a simple Markdown CV or cover letter into a clean Word (.docx) file.

Usage:
    python3 scripts/build_docx.py input.md output.docx [--letter]

Supported Markdown: # title, ## section, ### subheading, - bullets,
**bold**, *italic*, plain paragraphs, and --- (horizontal rule).
"""
import argparse
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def add_runs(paragraph, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3A5F")
    borders.append(bottom)
    p_pr.append(borders)


def build(md_text, letter):
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2 if letter else 1.6)
        section.left_margin = section.right_margin = Cm(2.2 if letter else 1.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11 if letter else 10.5)
    normal.paragraph_format.space_after = Pt(8 if letter else 3)
    normal.paragraph_format.line_spacing = 1.15 if letter else 1.05

    lines = md_text.splitlines()
    paragraph_buffer = []

    def flush():
        if paragraph_buffer:
            add_runs(doc.add_paragraph(), " ".join(paragraph_buffer))
            paragraph_buffer.clear()

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped == "---":
            flush()
            bottom_border(doc.add_paragraph())
            continue
        if stripped.startswith("# "):
            flush()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if letter else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.size = Pt(18 if letter else 20)
            run.font.color.rgb = ACCENT
            p.paragraph_format.space_after = Pt(2)
            continue
        if stripped.startswith("## "):
            flush()
            p = doc.add_paragraph()
            run = p.add_run(stripped[3:].upper())
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = ACCENT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            bottom_border(p)
            continue
        if stripped.startswith("### "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            add_runs(p, f"**{stripped[4:]}**" if "**" not in stripped else stripped[4:])
            continue
        if stripped.startswith(("- ", "* ", "• ")):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, stripped[2:])
            continue
        paragraph_buffer.append(stripped)

    flush()
    # A centred contact line right under a CV title reads better.
    if not letter and len(doc.paragraphs) > 1:
        doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--letter", action="store_true", help="cover letter layout")
    args = parser.parse_args()
    if not args.output.lower().endswith(".docx"):
        sys.exit("output must end in .docx")
    with open(args.input, encoding="utf-8") as f:
        doc = build(f.read(), args.letter)
    doc.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
